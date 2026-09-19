#!/usr/bin/env python3
"""Build a NON-DESTRUCTIVE migration manifest from Arham source workbooks."""
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook

def row_hash(values):
    payload = json.dumps([None if v is None else str(v) for v in values], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()

def audit(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        nonempty = 0; hashes = set(); duplicate_payloads = 0
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            if not any(v not in (None, "") for v in vals): continue
            nonempty += 1
            h = row_hash(vals)
            if h in hashes: duplicate_payloads += 1
            hashes.add(h)
        sheets.append({"name": ws.title, "nonempty_rows": nonempty, "duplicate_row_payloads": duplicate_payloads, "max_columns": ws.max_column})
    wb.close()
    return {"file": path.name, "sheets": sheets}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("-o", "--output", default="migration-manifest.json")
    args = ap.parse_args()
    manifest = {"migration_rules": {"non_destructive": True, "preserve_source_file_sheet_row": True, "stock_source_of_truth": "stock_movements", "current_stock_is_reconciliation_target": True, "costing_default": "weighted_average", "ambiguous_matches": "needs_review", "financial_stock_records": "never_hard_delete"}, "files": [audit(p) for p in args.files]}
    Path(args.output).write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Wrote {args.output}")

if __name__ == "__main__": main()
